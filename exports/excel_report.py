import pandas as pd
import os
import traceback
from engines.profiling_engine import profiling_engine
from engines.quality_engine import quality_engine
from engines.recommendation_engine import recommendation_engine
from datetime import datetime
from flask import request, send_file, redirect, flash
from openpyxl import load_workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.worksheet.table import Table as XLTable, TableStyleInfo


# ── BRAND PALETTE ────────────────────────────────────────────
XL_ACCENT      = "185FA5"   # brand blue
XL_ACCENT_DARK = "0C447C"
XL_ACCENT_LIGHT= "E6F1FB"
XL_BG          = "F5F4F0"   # off-white
XL_BORDER_CLR  = "E0DDD5"
XL_TEXT        = "111110"
XL_TEXT2       = "6B6A67"
XL_TEXT3       = "A09F9B"
XL_SUCCESS     = "2D6A11"
XL_SUCCESS_BG  = "EAF3DE"
XL_WARN        = "854F0B"
XL_WARN_BG     = "FAEEDA"
XL_DANGER      = "A32D2D"
XL_DANGER_BG   = "FCEBEB"
XL_WHITE       = "FFFFFF"


# ── STYLE HELPERS ────────────────────────────────────────────
def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, color=XL_TEXT, size=10, italic=False, name="Arial"):
    return Font(name=name, bold=bold, color=color, size=size, italic=italic)

def _align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def _border(color=XL_BORDER_CLR, style="thin"):
    s = Side(border_style=style, color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def _border_bottom(color=XL_BORDER_CLR):
    s = Side(border_style="thin", color=color)
    n = Side(border_style=None)
    return Border(left=n, right=n, top=n, bottom=s)

def _apply_header_row(ws, row_num, headers, col_start=1):
    """Apply branded blue header styling to a row."""
    for i, h in enumerate(headers, start=col_start):
        cell = ws.cell(row=row_num, column=i, value=h)
        cell.font      = _font(bold=True, color=XL_WHITE, size=9)
        cell.fill      = _fill(XL_ACCENT)
        cell.alignment = _align("center")
        cell.border    = _border(XL_ACCENT_DARK)

def _set_col_widths(ws, widths):
    """widths: list of (col_letter_or_index, width)"""
    for col, w in widths:
        if isinstance(col, int):
            col = get_column_letter(col)
        ws.column_dimensions[col].width = w

def _freeze(ws, cell="A2"):
    ws.freeze_panes = cell

def _autofit_col(ws, col_idx, min_w=10, max_w=40):
    col_letter = get_column_letter(col_idx)
    col_cells  = ws[col_letter]
    length = max(
        (len(str(c.value)) if c.value else 0) for c in col_cells
    )
    ws.column_dimensions[col_letter].width = min(max(length + 2, min_w), max_w)


# ── SHEET: COVER ─────────────────────────────────────────────
def _build_cover(wb, filename, quality_score, profile, df):
    ws = wb.create_sheet("Cover", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 20

    # ── Banner ──
    for r in range(1, 7):
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            cell.fill = _fill(XL_ACCENT)

    ws.row_dimensions[1].height = 6
    ws.row_dimensions[2].height = 36
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 16
    ws.row_dimensions[5].height = 16
    ws.row_dimensions[6].height = 10

    title_cell = ws.cell(row=2, column=2, value="CleanSight")
    title_cell.font      = _font(bold=True, color=XL_WHITE, size=24, name="Arial")
    title_cell.alignment = _align("left", "center")

    sub_cell = ws.cell(row=3, column=2, value="Data Quality Report")
    sub_cell.font      = _font(color="C8DFF5", size=12)
    sub_cell.alignment = _align("left", "center")

    ws.cell(row=4, column=2, value=f"Dataset: {filename}").font = _font(color="C8DFF5", size=9)
    ws.cell(row=5, column=2, value=f"Generated: {datetime.now().strftime('%d %b %Y, %H:%M')}").font = _font(color="C8DFF5", size=9)

    # ── Quality Score block ──
    ws.row_dimensions[8].height  = 14
    ws.row_dimensions[9].height  = 44
    ws.row_dimensions[10].height = 18
    ws.row_dimensions[11].height = 14

    if quality_score >= 80:
        q_color, q_bg, q_grade = XL_SUCCESS, XL_SUCCESS_BG, "Excellent"
    elif quality_score >= 60:
        q_color, q_bg, q_grade = XL_WARN, XL_WARN_BG, "Good"
    elif quality_score >= 40:
        q_color, q_bg, q_grade = XL_WARN, XL_WARN_BG, "Fair"
    else:
        q_color, q_bg, q_grade = XL_DANGER, XL_DANGER_BG, "Poor"

    ws.cell(row=8, column=2, value="QUALITY SCORE").font = _font(color=XL_TEXT3, size=8, bold=True)

    score_cell = ws.cell(row=9, column=2, value=f"{quality_score}/100")
    score_cell.font      = _font(bold=True, color=q_color, size=30)
    score_cell.alignment = _align("left", "center")
    score_cell.fill      = _fill(q_bg)
    for c in range(2, 5):
        ws.cell(row=9, column=c).fill = _fill(q_bg)

    grade_cell = ws.cell(row=10, column=2, value=q_grade)
    grade_cell.font      = _font(bold=True, color=q_color, size=11)
    grade_cell.alignment = _align("left", "center")

    # ── KPI Grid ──
    kpi_data = [
        ("Total Rows",    profile.get("rows", len(df)),            XL_ACCENT),
        ("Columns",       profile.get("columns", len(df.columns)), XL_ACCENT),
        ("Missing Vals",  profile.get("missing_values", ""),       XL_WARN),
        ("Duplicates",    profile.get("duplicate_rows", ""),       XL_DANGER),
        ("Numeric Cols",  profile.get("numeric_columns", ""),      XL_ACCENT),
        ("Text Cols",     profile.get("text_columns", ""),         XL_TEXT2),
    ]

    ws.row_dimensions[13].height = 14
    ws.row_dimensions[14].height = 36
    ws.row_dimensions[15].height = 18
    ws.row_dimensions[16].height = 6

    ws.cell(row=13, column=2, value="DATASET OVERVIEW").font = _font(color=XL_TEXT3, size=8, bold=True)

    # 6 KPI cells across columns B–G
    for i, (label, value, color) in enumerate(kpi_data):
        col = i + 2
        # Expand columns for KPIs
        ws.column_dimensions[get_column_letter(col)].width = 16

        val_cell = ws.cell(row=14, column=col, value=value)
        val_cell.font      = _font(bold=True, color=color, size=18)
        val_cell.fill      = _fill(XL_BG)
        val_cell.alignment = _align("center", "center")
        val_cell.border    = _border(XL_BORDER_CLR)

        lbl_cell = ws.cell(row=15, column=col, value=label)
        lbl_cell.font      = _font(color=XL_TEXT3, size=8)
        lbl_cell.fill      = _fill(XL_BG)
        lbl_cell.alignment = _align("center")
        lbl_cell.border    = _border_bottom(XL_BORDER_CLR)

    # ── Sheet index ──
    ws.row_dimensions[18].height = 14
    ws.cell(row=18, column=2, value="CONTENTS").font = _font(color=XL_TEXT3, size=8, bold=True)

    sheets_index = [
        ("Cleaned Data",       "Full cleaned dataset"),
        ("Missing Report",     "Missing value analysis with severity"),
        ("Profile Summary",    "Dataset profiling metrics"),
        ("AI Recommendations", "Automated cleaning recommendations"),
    ]

    for i, (sname, sdesc) in enumerate(sheets_index, start=19):
        ws.row_dimensions[i].height = 18
        n_cell = ws.cell(row=i, column=2, value=sname)
        n_cell.font      = _font(bold=True, color=XL_ACCENT, size=9)
        n_cell.alignment = _align("left", "center")
        n_cell.fill      = _fill(XL_ACCENT_LIGHT)
        n_cell.border    = _border(XL_BORDER_CLR)

        d_cell = ws.cell(row=i, column=3, value=sdesc)
        d_cell.font      = _font(color=XL_TEXT2, size=9, italic=True)
        d_cell.alignment = _align("left", "center")
        d_cell.fill      = _fill(XL_BG)
        d_cell.border    = _border(XL_BORDER_CLR)


# ── SHEET: CLEANED DATA ──────────────────────────────────────
def _style_cleaned_sheet(wb, df):
    ws = wb["Cleaned Data"]
    ws.sheet_view.showGridLines = False

    headers = list(df.columns)
    _apply_header_row(ws, 1, headers)
    _freeze(ws, "A2")

    for row_idx in range(2, len(df) + 2):
        bg = XL_WHITE if row_idx % 2 == 0 else XL_BG
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill      = _fill(bg)
            cell.font      = _font(size=9)
            cell.alignment = _align("left")
            cell.border    = _border_bottom(XL_BORDER_CLR)

    for col_idx in range(1, len(headers) + 1):
        _autofit_col(ws, col_idx)

    # Register as a table for Excel filter/sort UI
    if len(df) > 0:
        last_col = get_column_letter(len(headers))
        last_row = len(df) + 1
        tbl = XLTable(
            displayName="CleanedData",
            ref=f"A1:{last_col}{last_row}"
        )
        tbl_style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        tbl.tableStyleInfo = tbl_style
        ws.add_table(tbl)

    ws.row_dimensions[1].height = 22


# ── SHEET: MISSING REPORT ────────────────────────────────────
def _style_missing_sheet(wb, missing_report, total_rows):
    ws = wb["Missing Report"]
    ws.sheet_view.showGridLines = False

    # Title
    ws.row_dimensions[1].height = 28
    title = ws.cell(row=1, column=1, value="Missing Value Analysis")
    title.font      = _font(bold=True, color=XL_ACCENT, size=14)
    title.alignment = _align("left", "center")
    ws.merge_cells("A1:E1")

    # Subtitle
    ws.row_dimensions[2].height = 16
    sub = ws.cell(row=2, column=1, value=f"Total rows in dataset: {total_rows}")
    sub.font      = _font(color=XL_TEXT2, size=9, italic=True)
    sub.alignment = _align("left", "center")

    # Header
    ws.row_dimensions[3].height = 22
    headers = ["Column", "Missing Count", "Missing %", "Severity", "Visual"]
    _apply_header_row(ws, 3, headers)

    # Data rows
    for i, (col, count, pct) in enumerate(missing_report, start=4):
        ws.row_dimensions[i].height = 20
        pct_float = float(pct) if isinstance(pct, (int, float)) else float(str(pct).replace("%",""))

        if pct_float >= 50:
            sev, sev_color, row_bg = "HIGH",   XL_DANGER, XL_DANGER_BG
        elif pct_float >= 20:
            sev, sev_color, row_bg = "MEDIUM", XL_WARN,   XL_WARN_BG
        else:
            sev, sev_color, row_bg = "LOW",    XL_SUCCESS, XL_SUCCESS_BG

        # Column name
        c = ws.cell(row=i, column=1, value=col)
        c.font      = _font(bold=True, color=XL_ACCENT_DARK, size=9)
        c.fill      = _fill(XL_BG)
        c.alignment = _align("left")
        c.border    = _border(XL_BORDER_CLR)

        # Missing count
        c2 = ws.cell(row=i, column=2, value=count)
        c2.font      = _font(size=9)
        c2.fill      = _fill(XL_WHITE)
        c2.alignment = _align("center")
        c2.border    = _border(XL_BORDER_CLR)
        c2.number_format = "#,##0"

        # Percent
        c3 = ws.cell(row=i, column=3, value=pct_float / 100)
        c3.font          = _font(bold=True, color=sev_color, size=9)
        c3.fill          = _fill(XL_WHITE)
        c3.alignment     = _align("center")
        c3.border        = _border(XL_BORDER_CLR)
        c3.number_format = "0.00%"

        # Severity badge
        c4 = ws.cell(row=i, column=4, value=sev)
        c4.font      = _font(bold=True, color=XL_WHITE, size=8)
        c4.fill      = _fill(sev_color)
        c4.alignment = _align("center")
        c4.border    = _border(sev_color)

        # Visual bar (unicode approximation)
        bar_len  = int(pct_float / 5)
        bar_str  = "█" * bar_len + "░" * (20 - bar_len)
        c5 = ws.cell(row=i, column=5, value=bar_str)
        c5.font      = Font(name="Courier New", size=7, color=sev_color)
        c5.fill      = _fill(row_bg)
        c5.alignment = _align("left")
        c5.border    = _border(XL_BORDER_CLR)

    _set_col_widths(ws, [("A",22),("B",16),("C",14),("D",12),("E",26)])
    _freeze(ws, "A4")

    # Conditional color scale on % column
    if missing_report:
        last = 3 + len(missing_report)
        ws.conditional_formatting.add(
            f"C4:C{last}",
            ColorScaleRule(
                start_type="num", start_value=0,    start_color="EAF3DE",
                mid_type="num",   mid_value=0.25,   mid_color="FAEEDA",
                end_type="num",   end_value=1,       end_color="FCEBEB",
            )
        )

    # Bar chart of missing %
    if len(missing_report) >= 2:
        chart      = BarChart()
        chart.type = "bar"
        chart.title          = "Missing % by Column"
        chart.y_axis.title   = "Column"
        chart.x_axis.title   = "Missing %"
        chart.style          = 10
        chart.width          = 14
        chart.height         = max(6, len(missing_report) * 0.7)
        chart.grouping       = "clustered"
        chart.overlap        = 100

        data_ref = Reference(
            ws,
            min_col=3,
            min_row=3,
            max_row=3 + len(missing_report),
        )
        cats = Reference(
            ws,
            min_col=1,
            min_row=4,
            max_row=3 + len(missing_report),
        )
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats)
        chart.series[0].graphicalProperties.solidFill = XL_ACCENT

        ws.add_chart(chart, f"G3")


# ── SHEET: PROFILE SUMMARY ───────────────────────────────────
def _style_profile_sheet(wb, profile):
    ws = wb["Profile Summary"]
    ws.sheet_view.showGridLines = False

    ws.row_dimensions[1].height = 28
    title = ws.cell(row=1, column=1, value="Dataset Profile")
    title.font      = _font(bold=True, color=XL_ACCENT, size=14)
    title.alignment = _align("left", "center")
    ws.merge_cells("A1:C1")

    ws.row_dimensions[2].height = 8

    headers = ["Metric", "Value", "Notes"]
    _apply_header_row(ws, 3, headers)
    ws.row_dimensions[3].height = 22

    notes_map = {
        "rows":            "Total records in cleaned dataset",
        "columns":         "Number of fields / features",
        "missing_values":  "Total null cells across all columns",
        "duplicate_rows":  "Rows with identical values (removed)",
        "numeric_columns": "Columns with numeric data types",
        "text_columns":    "Columns with string / object data types",
    }

    for i, (metric, value) in enumerate(profile.items(), start=4):
        ws.row_dimensions[i].height = 20
        bg = XL_BG if i % 2 == 0 else XL_WHITE

        m = ws.cell(row=i, column=1, value=metric.replace("_", " ").title())
        m.font      = _font(bold=True, color=XL_TEXT, size=9)
        m.fill      = _fill(bg)
        m.alignment = _align("left")
        m.border    = _border(XL_BORDER_CLR)

        v = ws.cell(row=i, column=2, value=value)
        v.font      = _font(bold=True, color=XL_ACCENT, size=11)
        v.fill      = _fill(bg)
        v.alignment = _align("center")
        v.border    = _border(XL_BORDER_CLR)

        n = ws.cell(row=i, column=3, value=notes_map.get(metric, ""))
        n.font      = _font(color=XL_TEXT2, size=8, italic=True)
        n.fill      = _fill(bg)
        n.alignment = _align("left")
        n.border    = _border(XL_BORDER_CLR)

    _set_col_widths(ws, [("A",24), ("B",16), ("C",42)])
    _freeze(ws, "A4")


# ── SHEET: AI RECOMMENDATIONS ────────────────────────────────
def _style_reco_sheet(wb, recommendations):
    ws = wb["AI Recommendations"]
    ws.sheet_view.showGridLines = False

    ws.row_dimensions[1].height = 28
    title = ws.cell(row=1, column=1, value="AI Cleaning Recommendations")
    title.font      = _font(bold=True, color=XL_ACCENT, size=14)
    title.alignment = _align("left", "center")
    ws.merge_cells("A1:D1")

    ws.row_dimensions[2].height = 16
    sub = ws.cell(row=2, column=1,
        value="Auto-classified by CleanSight AI · Critical → Warning → Info")
    sub.font      = _font(color=XL_TEXT2, size=9, italic=True)
    sub.alignment = _align("left", "center")

    headers = ["#", "Recommendation", "Severity", "Category"]
    _apply_header_row(ws, 3, headers)
    ws.row_dimensions[3].height = 22

    category_map = {
        "email":     "Format",
        "phone":     "Format",
        "date":      "Format",
        "negative":  "Outlier",
        "outlier":   "Outlier",
        "missing":   "Missing",
        "duplicate": "Duplicate",
        "casing":    "Consistency",
        "inconsist": "Consistency",
        "invalid":   "Validation",
        "boolean":   "Format",
    }

    for i, rec in enumerate(recommendations, start=4):
        ws.row_dimensions[i].height = 22
        rec_lower = rec.lower()

        if any(k in rec_lower for k in ["invalid","error","negative","critical"]):
            sev, sev_color, bg = "CRITICAL", XL_DANGER, XL_DANGER_BG
        elif any(k in rec_lower for k in ["missing","inconsistent","mixed","duplicate"]):
            sev, sev_color, bg = "WARNING", XL_WARN, XL_WARN_BG
        else:
            sev, sev_color, bg = "INFO", XL_ACCENT, XL_ACCENT_LIGHT

        cat = next(
            (v for k, v in category_map.items() if k in rec_lower),
            "General"
        )

        num = ws.cell(row=i, column=1, value=i - 3)
        num.font      = _font(color=XL_TEXT3, size=8)
        num.fill      = _fill(bg)
        num.alignment = _align("center")
        num.border    = _border(XL_BORDER_CLR)

        txt = ws.cell(row=i, column=2, value=rec)
        txt.font      = _font(color=XL_TEXT, size=9)
        txt.fill      = _fill(bg)
        txt.alignment = _align("left", wrap=True)
        txt.border    = _border(XL_BORDER_CLR)

        sev_c = ws.cell(row=i, column=3, value=sev)
        sev_c.font      = _font(bold=True, color=XL_WHITE, size=8)
        sev_c.fill      = _fill(sev_color)
        sev_c.alignment = _align("center")
        sev_c.border    = _border(sev_color)

        cat_c = ws.cell(row=i, column=4, value=cat)
        cat_c.font      = _font(color=XL_ACCENT_DARK, size=8, bold=True)
        cat_c.fill      = _fill(XL_ACCENT_LIGHT)
        cat_c.alignment = _align("center")
        cat_c.border    = _border(XL_BORDER_CLR)

    _set_col_widths(ws, [("A",5), ("B",58), ("C",12), ("D",14)])
    _freeze(ws, "A4")

def generate_excel_report(filename,cleaned_folder,df,missing_report,profile,recommendations,quality_score):
    try:
        excel_path = os.path.join(
            cleaned_folder,
            "report_" + filename.replace(".csv", ".xlsx")
        )

        # ── Write raw data via pandas ──────────────────────
        with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Cleaned Data", index=False)

            pd.DataFrame(
                missing_report, columns=["Column","Missing Count","Percent"]
            ).to_excel(writer, sheet_name="Missing Report", index=False)

            pd.DataFrame(
                profile.items(), columns=["Metric","Value"]
            ).to_excel(writer, sheet_name="Profile Summary", index=False)

            pd.DataFrame(
                {"Recommendations": recommendations}
            ).to_excel(writer, sheet_name="AI Recommendations", index=False)

        # ── Re-open with openpyxl to apply full styling ────
        wb = load_workbook(excel_path)

        _build_cover(wb, filename, quality_score, profile, df)
        _style_cleaned_sheet(wb, df)
        _style_missing_sheet(wb, missing_report, len(df))
        _style_profile_sheet(wb, profile)
        _style_reco_sheet(wb, recommendations)

        # Set Cover as first active sheet
        wb.active = wb["Cover"]

        wb.save(excel_path)

        return excel_path
    except Exception:
        print(traceback.format_exc())
        return None